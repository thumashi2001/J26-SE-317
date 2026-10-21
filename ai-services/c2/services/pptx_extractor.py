
from pathlib import Path
from pptx import Presentation


def extract_text_from_pptx(file_path):
    """Extract slide text from a PPTX lecture presentation."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Presentation not found: {path}")

    if path.suffix.lower() != ".pptx":
        raise ValueError("Only PPTX files are supported")

    presentation = Presentation(str(path))
    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = []

        def collect_text(shapes):
            for shape in shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        text_parts.append(text)

                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            text = cell.text.strip()
                            if text:
                                text_parts.append(text)

                if shape.shape_type == 6:  # Grouped shapes
                    collect_text(shape.shapes)

        collect_text(slide.shapes)

        slides.append({
            "slide_number": slide_number,
            "text": "\n".join(text_parts),
        })

    full_text = "\n\n".join(
        slide["text"] for slide in slides if slide["text"]
    )

    return {
        "filename": path.name,
        "file_type": "pptx",
        "slide_count": len(slides),
        "character_count": len(full_text),
        "is_text_extractable": bool(full_text.strip()),
        "text": full_text,
        "slides": slides,
    }
